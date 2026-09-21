"""Matching JSON → GBP carga-masiva xlsx. USD without TC is RechazoExcel."""

from __future__ import annotations

import shutil
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TEMPLATE = Path(__file__).resolve().parent / "templates" / "carga_masiva_articulos.xlsx"

GBP_MONEDA_ARS = 1
INCLUIR_STOCK = 0
COSTO_ADICIONAL = 0
CENTAVO = Decimal("0.01")

COL_EAN = "Código/EAN"
COL_CANT = "Cantidad"
COL_STOCK = "Incluir en Stock Disponible"
COL_COSTO = "Costo"
COL_ADIC = "Costo Adicional"
COL_MONEDA = "Moneda"

FMT_EAN = "@"
FMT_COSTO = "0.000"
ANCHO_EAN = 22.0
CUIT_DISTECNA = "30714636827"


class RechazoExcel(Exception):
    """Do not generate the file (missing TC, unsupported currency, no ok rows)."""


def parse_tc(raw: object) -> Decimal | None:
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float, Decimal)):
        val = Decimal(str(raw))
    else:
        s = str(raw).strip().replace(" ", "")
        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            s = s.replace(",", ".")
        try:
            val = Decimal(s)
        except InvalidOperation:
            return None
    if val <= 0:
        return None
    return val


def moneda_norm(raw: object) -> str | None:
    u = str(raw or "").upper().replace("Ó", "O")
    u = u.replace("$", " ").strip()
    if not u:
        return None
    if "ARS" in u or "PESO" in u or u in {"ARG", "AR"}:
        return "ARS"
    if "USD" in u or "U$S" in u or "DOLAR" in u or "USS" in u:
        return "USD"
    return u[:12]


def cuit_digits(raw: object) -> str:
    return "".join(c for c in str(raw or "") if c.isdigit())


def es_distecna(matched: dict[str, Any]) -> bool:
    blobs = [matched.get("proveedor_razon_social"), matched.get("proveedor_cuit")]
    if any("DISTECNA" in str(b or "").upper() for b in blobs):
        return True
    return any(cuit_digits(b) == CUIT_DISTECNA for b in blobs)


def parse_descuento_frac(raw: object) -> Decimal | None:
    if raw is None or isinstance(raw, bool):
        return None
    try:
        val = Decimal(str(raw).strip().replace(",", ".").replace("%", ""))
    except InvalidOperation:
        return None
    if val <= 0:
        return None
    if val > 1:
        val = val / Decimal(100)
    if val >= 1:
        return None
    return val


def tipo_cambio_oc(matched: dict[str, Any]) -> tuple[Decimal | None, Decimal | None, Decimal | None]:
    """(tc_papel, descuento_frac, tc_para_oc). No fuente-file lookup — Pricing SoT is matched."""
    tc = parse_tc(matched.get("tipo_cambio"))
    dto = parse_descuento_frac(matched.get("descuento_pct"))
    if tc is None:
        return None, dto, None
    if dto and es_distecna(matched):
        return tc, dto, (tc / (Decimal(1) - dto))
    return tc, dto, tc


def costo_ars(renglon: dict[str, Any], tc: Decimal | None, moneda_doc: str | None) -> Decimal:
    precio = renglon.get("precio_unitario")
    if precio is None:
        raise RechazoExcel("renglón ok sin precio_unitario")
    try:
        p = Decimal(str(precio))
    except InvalidOperation as exc:
        raise RechazoExcel(f"precio_unitario inválido: {precio}") from exc
    mon = moneda_norm(renglon.get("moneda")) or moneda_norm(moneda_doc)
    if mon == "ARS":
        return p.quantize(CENTAVO, rounding=ROUND_HALF_UP)
    if mon == "USD":
        if tc is None:
            raise RechazoExcel(
                "hay renglones en USD y no hay tipo de cambio en el adjunto; se rechaza (no se inventa TC)"
            )
        return (p * tc).quantize(CENTAVO, rounding=ROUND_HALF_UP)
    if mon is None:
        raise RechazoExcel("renglón ok sin moneda")
    raise RechazoExcel(f"moneda no soportada para OC: {mon}")


def filas_ok(matched: dict[str, Any], tc: Decimal | None) -> list[dict[str, Any]]:
    moneda_doc = matched.get("moneda")
    out: list[dict[str, Any]] = []
    for r in matched.get("renglones") or []:
        m = r.get("match") or {}
        if m.get("estado") != "ok":
            continue
        ean = str(m.get("ean") or "").strip()
        if not ean:
            raise RechazoExcel("renglón ok sin EAN de GBP")
        cant = r.get("cantidad")
        if cant is None:
            raise RechazoExcel(f"renglón ok {ean} sin cantidad")
        try:
            cantidad = int(cant)
        except (TypeError, ValueError) as exc:
            raise RechazoExcel(f"cantidad inválida en {ean}: {cant}") from exc
        if cantidad <= 0:
            raise RechazoExcel(f"cantidad no positiva en {ean}: {cantidad}")
        costo = costo_ars(r, tc, moneda_doc)
        out.append(
            {
                COL_EAN: ean,
                COL_CANT: cantidad,
                COL_STOCK: int(INCLUIR_STOCK),
                COL_COSTO: float(costo),
                COL_ADIC: int(COSTO_ADICIONAL),
                COL_MONEDA: int(GBP_MONEDA_ARS),
            }
        )
    return out


def escribir_xlsx(filas: list[dict[str, Any]], dest: Path) -> None:
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"No está la plantilla: {TEMPLATE}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(TEMPLATE, dest)
    wb = load_workbook(dest)
    ws = wb["Artículos"]
    headers = [ws.cell(1, c).value for c in range(1, (ws.max_column or 1) + 1)]
    col_idx = {name: i + 1 for i, name in enumerate(headers) if name}
    missing = [n for n in (COL_EAN, COL_CANT, COL_STOCK, COL_COSTO, COL_ADIC, COL_MONEDA) if n not in col_idx]
    if missing:
        wb.close()
        dest.unlink(missing_ok=True)
        raise RechazoExcel(f"plantilla: faltan columnas {missing}")
    col_ean = col_idx[COL_EAN]
    dim_ean = ws.column_dimensions[get_column_letter(col_ean)]
    if dim_ean.width is None or dim_ean.width < ANCHO_EAN:
        dim_ean.width = ANCHO_EAN
    for r_i, fila in enumerate(filas, 2):
        c_ean = ws.cell(r_i, col_ean, str(fila[COL_EAN]))
        c_ean.number_format = FMT_EAN
        c_ean.data_type = "s"
        ws.cell(r_i, col_idx[COL_CANT], int(fila[COL_CANT]))
        ws.cell(r_i, col_idx[COL_STOCK], int(fila[COL_STOCK]))
        c_costo = ws.cell(r_i, col_idx[COL_COSTO], float(fila[COL_COSTO]))
        c_costo.number_format = FMT_COSTO
        ws.cell(r_i, col_idx[COL_ADIC], int(fila[COL_ADIC]))
        ws.cell(r_i, col_idx[COL_MONEDA], int(fila[COL_MONEDA]))
    wb.active = ws
    wb.save(dest)
    wb.close()


def generar(matched: dict[str, Any], dest: Path) -> dict[str, Any]:
    ok_n = sum(1 for r in (matched.get("renglones") or []) if (r.get("match") or {}).get("estado") == "ok")
    if ok_n == 0:
        raise RechazoExcel("no hay renglones ok; no se genera Excel")
    tc_papel, dto, tc = tipo_cambio_oc(matched)
    filas = filas_ok(matched, tc)
    if not filas:
        raise RechazoExcel("no hay renglones ok; no se genera Excel")
    escribir_xlsx(filas, dest)
    return {
        "archivo": dest.name,
        "filas": len(filas),
        "tipo_cambio": str(tc) if tc else None,
        "tipo_cambio_papel": str(tc_papel) if tc_papel else None,
        "descuento_pct": str(dto * 100) if dto else None,
        "moneda_gbp": GBP_MONEDA_ARS,
    }
