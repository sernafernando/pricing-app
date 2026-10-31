"""Build Gestión Full Excel exports (official ML template + ops EAN sheet)."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from sqlalchemy import bindparam, text
from sqlalchemy.orm import Session

from app.models.gestion_full_package import (
    PACKAGE_STATUS_EXPORTED,
    PACKAGE_STATUS_LOCKED,
    PACKAGE_STATUS_OPEN,
    GestionFullPackage,
    GestionFullPackageLine,
)
from app.models.gestion_full_unidades_caja import GestionFullUnidadesCaja

_TEMPLATE = (
    Path(__file__).resolve().parents[2]
    / "assets"
    / "gestion_full"
    / "envio-de-stock-full.xlsx"
)
_SHEET = " Selección de productos"
_DATA_START_ROW = 6

_ENRICH_SQL = text(
    """
    SELECT DISTINCT ON (i.inventory_id)
      i.inventory_id,
      i.seller_sku,
      i.item_id AS mla
    FROM ml_items i
    WHERE i.inventory_id IN :ids
      AND i.never_existed IS NOT TRUE
      AND i.gone_at IS NULL
    ORDER BY i.inventory_id, i.item_id
    """
).bindparams(bindparam("ids", expanding=True))

_EAN_SQL = text(
    """
    SELECT p.item_id AS producto_item_id, p.codigo AS ean, p.descripcion
    FROM productos_erp p
    WHERE p.item_id IN :ids
    """
).bindparams(bindparam("ids", expanding=True))


def _mla_digits(mla: Optional[str]) -> Optional[str]:
    if not mla:
        return None
    s = str(mla).strip().upper()
    if s.startswith("MLA"):
        s = s[3:]
    return s if s.isdigit() else None


def _enrich_lines(
    db: Session,
    lines: list[GestionFullPackageLine],
) -> list[dict[str, Any]]:
    inv_ids = [ln.inventory_id for ln in lines]
    erp_ids = [ln.producto_item_id for ln in lines]
    by_inv: dict[str, Any] = {}
    if inv_ids:
        for row in db.execute(_ENRICH_SQL, {"ids": inv_ids}).mappings():
            by_inv[row.inventory_id] = row
    by_erp: dict[int, Any] = {}
    if erp_ids:
        for row in db.execute(_EAN_SQL, {"ids": erp_ids}).mappings():
            by_erp[row.producto_item_id] = row

    out: list[dict[str, Any]] = []
    for ln in lines:
        inv = by_inv.get(ln.inventory_id)
        erp = by_erp.get(ln.producto_item_id)
        out.append(
            {
                "line": ln,
                "seller_sku": (inv.seller_sku if inv else None) or (erp.ean if erp else None),
                "mla_digits": _mla_digits(inv.mla if inv else None),
                "ean": erp.ean if erp else None,
                "descripcion": erp.descripcion if erp else None,
            }
        )
    return out


def assert_package_exportable(db: Session, package: GestionFullPackage) -> None:
    """Raise ValueError when package cannot be exported."""
    if package.status == PACKAGE_STATUS_LOCKED:
        raise ValueError("paquete hard-locked")
    if any(ln.inbound_id for ln in package.lines):
        raise ValueError("hay inbound_id asignados; borrarlos antes de regenerar")
    if not package.lines:
        raise ValueError("paquete sin líneas")
    missing_caja: list[int] = []
    for ln in package.lines:
        if ln.qty_export < 1:
            continue
        caja = db.get(GestionFullUnidadesCaja, ln.producto_item_id)
        if caja is None:
            missing_caja.append(ln.producto_item_id)
    if missing_caja:
        raise ValueError(
            f"faltan unidades/caja en productos: {sorted(set(missing_caja))}"
        )


def build_ml_full_xlsx(db: Session, package: GestionFullPackage) -> bytes:
    """Fill the official ML Full template from package qty_export (≥1)."""
    assert_package_exportable(db, package)
    if not _TEMPLATE.is_file():
        raise FileNotFoundError(f"template missing: {_TEMPLATE}")

    enriched = _enrich_lines(db, list(package.lines))
    rows = [
        e
        for e in enriched
        if e["line"].qty_export >= 1 and e["seller_sku"] and e["mla_digits"]
    ]
    if not rows:
        raise ValueError("ninguna línea exportable (SKU + MLA + qty≥1)")

    wb = load_workbook(io.BytesIO(_TEMPLATE.read_bytes()))
    ws = wb[_SHEET]
    for offset, e in enumerate(rows):
        r = _DATA_START_ROW + offset
        ln = e["line"]
        ws.cell(r, 1, value=str(e["seller_sku"])).number_format = "@"
        ws.cell(r, 3, value=str(ln.inventory_id)).number_format = "@"
        ws.cell(r, 4, value=str(e["mla_digits"])).number_format = "@"
        ws.cell(r, 6, value=str(int(ln.qty_export))).number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_ops_ean_xlsx(db: Session, package: GestionFullPackage) -> bytes:
    """Simple ops sheet: EAN, descripción, cantidad (qty_export)."""
    assert_package_exportable(db, package)
    enriched = _enrich_lines(db, list(package.lines))
    wb = Workbook()
    ws = wb.active
    ws.title = "Envio"
    ws.append(["EAN", "Descripcion", "Cantidad"])
    for e in enriched:
        ln = e["line"]
        if ln.qty_export < 1:
            continue
        ws.append([e["ean"] or "", e["descripcion"] or "", int(ln.qty_export)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mark_package_exported(db: Session, package: GestionFullPackage) -> GestionFullPackage:
    """Flip open → exported on successful download (idempotent if already exported)."""
    if package.status == PACKAGE_STATUS_OPEN:
        package.status = PACKAGE_STATUS_EXPORTED
        package.exported_at = datetime.now(UTC)
        db.commit()
        db.refresh(package)
    return package
