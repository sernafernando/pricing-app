"""
The four price-list exports and the current-view export of Productos, as Publicaciones ML reuses them (S12c):

- Each kind checks its own `productos.exportar_*` permission on the server, as the export modal does in the UI:
  Rebate ML, Web Transferencia, Clásica, PVP (the Clásica route with a `pvp*` `tipo_cuotas`) and Vista actual.
- An optional explicit product set (`item_ids`, a JSON body: POST on every route, the Rebate body for Rebate) exports
  exactly those products -- the set Publicaciones ML's filters select, resolved by `/ml-publications/view/product-ids`.
  It combines with the query filters; an empty list or one over the cap is a 422; without it nothing changes.

Driven end to end through the HTTP client, reading the exported workbook back.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.security import get_password_hash
from app.models.producto import ProductoERP, ProductoPricing
from app.models.precio_ml import PrecioML
from app.models.publicacion_ml import PublicacionML
from app.models.usuario import AuthProvider, RolUsuario, Usuario
from app.services.ml_publications.view.caps import MAX_EXPLICIT_PRODUCT_IDS

PERMISOS = {
    "vista_actual": "productos.exportar_vista_actual",
    "rebate": "productos.exportar_rebate",
    "web_transferencia": "productos.exportar_web_transferencia",
    "clasica": "productos.exportar_clasica",
    "pvp": "productos.exportar_pvp",
}

# (kind, method, path, request kwargs) -- the five exports as the modal asks for them.
EXPORTS = [
    pytest.param("vista_actual", "get", "/api/exportar-vista-actual", {}, id="vista_actual"),
    pytest.param("rebate", "post", "/api/productos/exportar-rebate", {"json": {}}, id="rebate"),
    pytest.param("web_transferencia", "get", "/api/exportar-web-transferencia", {}, id="web_transferencia"),
    pytest.param("clasica", "get", "/api/exportar-clasica?tipo_cuotas=clasica", {}, id="clasica"),
    pytest.param("pvp", "get", "/api/exportar-clasica?tipo_cuotas=pvp_6", {}, id="pvp"),
]


def _headers(user: Usuario) -> dict:
    from tests.conftest import make_access_token

    return {"Authorization": f"Bearer {make_access_token(user)}"}


@pytest.fixture()
def usuario(db, rol_ventas) -> Usuario:
    user = Usuario(
        username="exportador",
        email="exportador@example.com",
        nombre="exportador",
        password_hash=get_password_hash("TestPass123!"),
        rol=RolUsuario.VENTAS,
        rol_id=rol_ventas.id,
        auth_provider=AuthProvider.LOCAL,
        activo=True,
    )
    db.add(user)
    db.commit()
    return user


@pytest.fixture()
def productos(db) -> None:
    """Three products that every export lists: priced on every channel, in rebate, with a publication."""
    for item_id in (1, 2, 3):
        db.add(
            ProductoERP(
                item_id=item_id,
                codigo=f"COD{item_id}",
                descripcion=f"Producto {item_id}",
                costo=100.0,
                iva=21.0,
                stock=5,
                marca="MARCA",
                activo=True,
            )
        )
        db.add(
            ProductoPricing(
                item_id=item_id,
                precio_lista_ml=1000.0,
                markup_calculado=10.0,
                participa_rebate=True,
                porcentaje_rebate=3.8,
                participa_web_transferencia=True,
                precio_web_transferencia=1200.0,
                precio_pvp_6_cuotas=1500.0,
            )
        )
        db.add(PrecioML(item_id=item_id, pricelist_id=4, precio=1100.0))
        db.add(PublicacionML(mla=f"MLA{item_id}", item_id=item_id, activo=True, pricelist_id=4))
    db.commit()


def _codigos(content: bytes) -> set[str]:
    """The products of the workbook as `COD<n>`: each kind puts the code in its own column, and Rebate (one row per
    publication) only the MLA, `MLA<n>` for product n here."""
    ws = load_workbook(BytesIO(content)).active
    cells = {str(v) for row in ws.iter_rows(min_row=2, values_only=True) for v in row if v is not None}
    return {c for c in cells if c.startswith("COD")} | {"COD" + c[3:] for c in cells if c.startswith("MLA")}


def _request(client, user, method: str, path: str, kwargs: dict, item_ids=None):
    kwargs = {**kwargs}
    if item_ids is not None:
        if method == "get":
            method = "post"  # the explicit set travels in a body: a long id list does not fit a URL
            kwargs["json"] = {"item_ids": item_ids}
        else:
            kwargs["json"] = {**kwargs.get("json", {}), "item_ids": item_ids}
    return getattr(client, method)(path, headers=_headers(user), **kwargs)


class TestPermisoPorTipo:
    @pytest.mark.parametrize("kind,method,path,kwargs", EXPORTS)
    def test_without_its_permission_is_403(
        self, client, db, usuario, grant_permiso, rol_ventas, kind, method, path, kwargs
    ) -> None:
        for other, codigo in PERMISOS.items():
            if other != kind:
                grant_permiso(rol_ventas, codigo)
        db.commit()

        resp = _request(client, usuario, method, path, kwargs)

        assert resp.status_code == 403

    @pytest.mark.parametrize("kind,method,path,kwargs", EXPORTS)
    def test_with_its_permission_exports(
        self, client, db, usuario, productos, grant_permiso, rol_ventas, kind, method, path, kwargs
    ) -> None:
        grant_permiso(rol_ventas, PERMISOS[kind])
        db.commit()

        resp = _request(client, usuario, method, path, kwargs)

        assert resp.status_code == 200, resp.text
        assert _codigos(resp.content) == {"COD1", "COD2", "COD3"}


class TestConjuntoExplicito:
    @pytest.fixture(autouse=True)
    def _todos_los_permisos(self, db, grant_permiso, rol_ventas) -> None:
        for codigo in PERMISOS.values():
            grant_permiso(rol_ventas, codigo)
        db.commit()

    @pytest.mark.parametrize("kind,method,path,kwargs", EXPORTS)
    def test_exports_exactly_the_given_products(self, client, usuario, productos, kind, method, path, kwargs) -> None:
        resp = _request(client, usuario, method, path, kwargs, item_ids=[3, 1, 3])

        assert resp.status_code == 200, resp.text
        assert _codigos(resp.content) == {"COD1", "COD3"}

    def test_combines_with_the_query_filters(self, client, usuario, productos) -> None:
        resp = _request(client, usuario, "get", "/api/exportar-vista-actual?search=COD3", {}, item_ids=[1, 3])

        assert resp.status_code == 200, resp.text
        assert _codigos(resp.content) == {"COD3"}

    def test_rebate_refuses_filtros_alongside(self, client, usuario, productos) -> None:
        resp = _request(
            client,
            usuario,
            "post",
            "/api/productos/exportar-rebate",
            {"json": {"filtros": {"search": "x"}}},
            item_ids=[1],
        )

        assert resp.status_code == 422

    @pytest.mark.parametrize("kind,method,path,kwargs", EXPORTS)
    def test_empty_list_is_422_never_all(self, client, usuario, productos, kind, method, path, kwargs) -> None:
        resp = _request(client, usuario, method, path, kwargs, item_ids=[])

        assert resp.status_code == 422

    @pytest.mark.parametrize("kind,method,path,kwargs", EXPORTS)
    def test_over_the_cap_is_422(self, client, usuario, productos, kind, method, path, kwargs) -> None:
        resp = _request(client, usuario, method, path, kwargs, item_ids=list(range(1, MAX_EXPLICIT_PRODUCT_IDS + 2)))

        assert resp.status_code == 422

    @pytest.mark.parametrize("kind,method,path,kwargs", EXPORTS)
    def test_without_it_nothing_changes(self, client, usuario, productos, kind, method, path, kwargs) -> None:
        resp = _request(client, usuario, method, path, kwargs)

        assert resp.status_code == 200, resp.text
        assert _codigos(resp.content) == {"COD1", "COD2", "COD3"}
